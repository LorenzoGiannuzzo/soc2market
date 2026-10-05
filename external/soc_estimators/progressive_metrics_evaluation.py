"""
================================================================================
SOC Prediction Framework for LFP Battery Cells
================================================================================

This work was originally developed in the context of the Huawei Tech Arena
Challenge (Nuremberg, 2024) for State of Charge (SOC) prediction of LFP
(Lithium Iron Phosphate) battery cells, where it was awarded First Prize.

The framework was subsequently extended and refined by researchers at the
Energy Center Lab, Department of Energy (DENERG), Politecnico di Torino.
The codebase, particularly the AI/Machine Learning model implementations,
was primarily developed by Lorenzo Giannuzzo.

Authors:
    Lorenzo Giannuzzo        - lorenzo.giannuzzo@polito.it
    Elena Rozzi               - elena.rozzi@polito.it
    Marcel Stolte             - marcel.stolte@polito.it
    Daniele Salvatore Schiera - daniele.schiera@polito.it
    Francesco D. Minuto       - francesco.minuto@polito.it
    Andrea Lanzini            - andrea.lanzini@polito.it

Supervision:
    Mansour Alramlawi         - Huawei Technologies
    (supervision of the related research paper and selected code modules)

Affiliation:
    Politecnico di Torino
    Department of Energy (DENERG) - Energy Center Lab
    Corso Duca degli Abruzzi 24, 10129 Torino, Italy

================================================================================
"""

"""
================================================================================
Progressive Temporal Window Metrics Evaluation
================================================================================

Reads pre-computed prediction results from results/results_sheet/ and evaluates
RMSE, MeanAE, MaxAE at progressive time windows (6h, 12h, 24h).

The results_sheet folder contains one Excel file per model (ANN.xlsx, CNN.xlsx, etc.)
with one sheet per scenario containing SOC_true and Predicted_SOC columns.

Output: Excel file with progressive metrics comparison across all models and scenarios.

================================================================================
"""

import numpy as np
import pandas as pd
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ============================================================================
# CONFIGURATION
# ============================================================================
SECONDS_PER_HOUR = 3600
MAX_HOURS = 24                          # Maximum evaluation window (hours)
HOURS = [6, 12, 24]                     # Time windows in hours
HOURS_IN_SECONDS = {h: h * SECONDS_PER_HOUR for h in HOURS}  # {6: 21600, 12: 43200, 24: 86400}
METRICS = ['RMSE', 'MeanAE', 'MaxAE']

MODEL_NAMES = ['ANN', 'CNN', 'LSTM', 'CatBoost', 'TFT', 'Ensemble']
MODEL_COLORS = {
    'ANN':      'D6E4F0',  # Light blue
    'CNN':      'FDE9D9',  # Light orange
    'LSTM':     'E4DFEC',  # Light purple
    'CatBoost': 'D8E4BC',  # Light green
    'TFT':      'FCE4D6',  # Light peach
    'Ensemble': 'C6EFCE',  # Light mint green
}

# Column names in results files
TRUE_COL = 'SOC_true'
PRED_COL = 'Predicted_SOC'

# Paths
SCRIPT_DIR = Path(__file__).resolve().parent
RESULTS_SHEET_DIR = SCRIPT_DIR / 'results' / 'results_sheet'
OUTPUT_DIR = SCRIPT_DIR / 'results'


# ============================================================================
# METRICS
# ============================================================================
def calc_rmse(y_true, y_pred):
    return np.sqrt(np.mean((y_true - y_pred) ** 2))

def calc_mean_ae(y_true, y_pred):
    return np.mean(np.abs(y_true - y_pred))

def calc_max_ae(y_true, y_pred):
    return np.max(np.abs(y_true - y_pred))


# ============================================================================
# EXCEL CREATION
# ============================================================================
def create_excel_report(all_results, output_path):
    """
    Create a formatted Excel report with progressive metrics.
    Single sheet with Dataset and Time Window columns.

    all_results: dict of {scenario_name: {model_name: {hours: {metric: value}}}}
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Progressive Metrics"

    # Styles
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )
    header_font = Font(name='Arial', bold=True, size=11, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1F4E79')
    data_font = Font(name='Arial', size=10)
    center = Alignment(horizontal='center', vertical='center')
    left_align = Alignment(horizontal='left', vertical='center')

    available_models = [m for m in MODEL_NAMES if any(
        m in scenario_data for scenario_data in all_results.values()
    )]

    # ============================================================
    # Build header row
    # ============================================================
    headers = ['Dataset', 'Time Window']
    for model in available_models:
        for metric in METRICS:
            headers.append(f'{model}_{metric}')

    # Write headers
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = thin_border

    # ============================================================
    # Write data rows
    # ============================================================
    row = 2
    for scenario_name in sorted(all_results.keys()):
        scenario_data = all_results[scenario_name]

        for h in HOURS:
            # Dataset name
            cell = ws.cell(row=row, column=1, value=scenario_name)
            cell.font = data_font
            cell.alignment = left_align
            cell.border = thin_border

            # Time window label in hours (e.g. "6h", "12h", "24h")
            cell = ws.cell(row=row, column=2, value=f'{h}h')
            cell.font = data_font
            cell.alignment = center
            cell.border = thin_border

            # Metrics for each model
            col = 3
            for model_name in available_models:
                model_metrics = scenario_data.get(model_name, {}).get(h, {})
                color = MODEL_COLORS.get(model_name, 'FFFFFF')
                model_fill = PatternFill('solid', fgColor=color)

                for metric in METRICS:
                    val = model_metrics.get(metric, None)
                    cell = ws.cell(row=row, column=col)
                    if val is not None:
                        cell.value = round(val, 4)
                    else:
                        cell.value = 'N/A'
                    cell.font = data_font
                    cell.alignment = center
                    cell.border = thin_border
                    cell.fill = model_fill
                    col += 1

            row += 1

    # ============================================================
    # Add AVERAGE row at the bottom
    # ============================================================
    row += 1  # Empty row before average

    for h in HOURS:
        # Dataset name = "AVERAGE"
        cell = ws.cell(row=row, column=1, value="AVERAGE")
        cell.font = Font(name='Arial', bold=True, size=10)
        cell.alignment = left_align
        cell.border = thin_border
        cell.fill = PatternFill('solid', fgColor='FFC000')  # Gold for average

        # Time window label
        cell = ws.cell(row=row, column=2, value=f'{h}h')
        cell.font = Font(name='Arial', bold=True, size=10)
        cell.alignment = center
        cell.border = thin_border
        cell.fill = PatternFill('solid', fgColor='FFC000')

        # Average metrics for each model
        col = 3
        for model_name in available_models:
            color = MODEL_COLORS.get(model_name, 'FFFFFF')

            for metric in METRICS:
                values = []
                for scenario_data in all_results.values():
                    if model_name in scenario_data and h in scenario_data[model_name]:
                        v = scenario_data[model_name][h].get(metric)
                        if v is not None:
                            values.append(v)

                cell = ws.cell(row=row, column=col)
                if values:
                    cell.value = round(np.mean(values), 4)
                else:
                    cell.value = 'N/A'
                cell.font = Font(name='Arial', bold=True, size=10)
                cell.alignment = center
                cell.border = thin_border
                cell.fill = PatternFill('solid', fgColor='FFC000')
                col += 1

        row += 1

    # ============================================================
    # Column widths
    # ============================================================
    ws.column_dimensions['A'].width = 25  # Dataset name
    ws.column_dimensions['B'].width = 14  # Time Window
    for c in range(3, 3 + len(available_models) * len(METRICS)):
        ws.column_dimensions[get_column_letter(c)].width = 14

    # Freeze header row
    ws.freeze_panes = 'A2'

    wb.save(str(output_path))
    print(f"\n  Excel saved: {output_path}")


# ============================================================================
# MAIN
# ============================================================================
def main():
    OUTPUT_DIR.mkdir(exist_ok=True)

    print("=" * 70)
    print("  PROGRESSIVE TEMPORAL WINDOW METRICS EVALUATION")
    print("=" * 70)
    print(f"  Results folder: {RESULTS_SHEET_DIR}")
    print(f"  Models:         {', '.join(MODEL_NAMES)}")
    print(f"  Max window:     {MAX_HOURS}h ({MAX_HOURS * SECONDS_PER_HOUR:,} samples)")
    print(f"  Windows:        {', '.join(f'{h}h ({HOURS_IN_SECONDS[h]:,}s)' for h in HOURS)}")
    print(f"  Metrics:        {', '.join(METRICS)}")
    print("=" * 70)

    # Check if results_sheet folder exists
    if not RESULTS_SHEET_DIR.exists():
        print(f"\nERROR: Results folder not found: {RESULTS_SHEET_DIR}")
        print("Run main.py first to generate prediction results.")
        return

    # Collect all scenarios from all models
    all_scenarios = set()
    model_data = {}  # {model_name: {scenario_name: DataFrame}}

    for model_name in MODEL_NAMES:
        model_file = RESULTS_SHEET_DIR / f'{model_name}.xlsx'
        if not model_file.exists():
            print(f"\n  [SKIP] {model_name}.xlsx not found")
            continue

        print(f"\n  Loading {model_name}.xlsx...")

        try:
            xl = pd.ExcelFile(model_file)
        except Exception as e:
            print(f"    [ERROR] Could not open file: {e}")
            print(f"    File may be corrupted. Re-run main.py to regenerate.")
            continue

        model_data[model_name] = {}

        for sheet_name in xl.sheet_names:
            if sheet_name == 'Metrics':
                continue  # Skip the metrics summary sheet

            try:
                df = pd.read_excel(xl, sheet_name=sheet_name)
            except Exception as e:
                print(f"    [ERROR] Could not read sheet '{sheet_name}': {e}")
                continue

            # Check required columns
            if TRUE_COL not in df.columns or PRED_COL not in df.columns:
                print(f"    [SKIP] {sheet_name}: missing {TRUE_COL} or {PRED_COL}")
                continue

            all_scenarios.add(sheet_name)
            model_data[model_name][sheet_name] = df

        print(f"    Loaded {len(model_data[model_name])} scenarios")

    if not model_data:
        print("\nERROR: No model data loaded. Check that results files exist.")
        return

    print(f"\n  Total unique scenarios: {len(all_scenarios)}")

    # Calculate metrics for each scenario, model, and time window
    all_results = {}  # {scenario: {model: {hours: {metric: value}}}}

    max_samples = MAX_HOURS * SECONDS_PER_HOUR  # Hard cap at 24h

    for scenario in sorted(all_scenarios):
        print(f"\n  Processing: {scenario}")
        all_results[scenario] = {}

        for model_name in MODEL_NAMES:
            if model_name not in model_data or scenario not in model_data[model_name]:
                continue

            df = model_data[model_name][scenario]
            y_true = df[TRUE_COL].values.astype(float)
            y_pred = df[PRED_COL].values.astype(float)

            # Remove NaN values
            valid = ~np.isnan(y_true) & ~np.isnan(y_pred)
            y_true = y_true[valid]
            y_pred = y_pred[valid]

            # Cap to 24h maximum
            y_true = y_true[:max_samples]
            y_pred = y_pred[:max_samples]

            model_metrics = {}
            for h in HOURS:
                end_idx = HOURS_IN_SECONDS[h]
                end_idx = min(end_idx, len(y_true))  # Don't exceed actual data

                if end_idx < 2:
                    model_metrics[h] = {m: None for m in METRICS}
                    continue

                yt = y_true[:end_idx]
                yp = y_pred[:end_idx]

                model_metrics[h] = {
                    'RMSE':   calc_rmse(yt, yp),
                    'MeanAE': calc_mean_ae(yt, yp),
                    'MaxAE':  calc_max_ae(yt, yp),
                }

            all_results[scenario][model_name] = model_metrics

            # Print 24h RMSE
            rmse_24h = model_metrics.get(24, {}).get('RMSE', 'N/A')
            if isinstance(rmse_24h, float):
                print(f"    {model_name}: RMSE@24h={rmse_24h:.4f}")

    # Create Excel report
    output_path = OUTPUT_DIR / 'progressive_metrics_comparison.xlsx'
    create_excel_report(all_results, output_path)

    print("\n" + "=" * 70)
    print("  DONE")
    print("=" * 70)


if __name__ == "__main__":
    main()